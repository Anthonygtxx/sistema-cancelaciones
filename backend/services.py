import os
import fitz  # PyMuPDF
import re
from docx import Document

# Constante para calcular Veces Salario Mínimo Mensual en DF/CDMX
SALARIO_MINIMO_MENSUAL_DF = 3713.40

def limpiar_ceros_izquierda(val):
    val = str(val).strip()
    if val.isdigit():
        return str(int(val))
    val_limpia = val.lstrip('0')
    return val_limpia if val_limpia else "0"

def credito_a_letras(numero_str):
    mapa_digitos = {
        '0': 'cero', '1': 'uno', '2': 'dos', '3': 'tres', '4': 'cuatro',
        '5': 'cinco', '6': 'seis', '7': 'siete', '8': 'ocho', '9': 'nueve'
    }
    digitos = [mapa_digitos[d] for d in str(numero_str).strip() if d in mapa_digitos]
    if digitos:
        return f'"{numero_str}" ({" ".join(digitos)})'
    return str(numero_str)

def limpiar_fecha_escritura(fecha_str):
    if not fecha_str or fecha_str == "NO_ENCONTRADO":
        return ""
    
    romanos = {
        'I': '01', 'II': '02', 'III': '03', 'IV': '04', 'V': '05', 'VI': '06',
        'VII': '07', 'VIII': '08', 'IX': '09', 'X': '10', 'XI': '11', 'XII': '12'
    }
    meses_texto = {
        '01': 'enero', '02': 'febrero', '03': 'marzo', '04': 'abril', '05': 'mayo', '06': 'junio',
        '07': 'julio', '08': 'agosto', '09': 'septiembre', '10': 'octubre', '11': 'noviembre', '12': 'diciembre'
    }
    
    m = re.search(r'(\d{1,2})[-/]([I|V|X]+)[-/](\d{4})', fecha_str, re.IGNORECASE)
    if m:
        dia = m.group(1)
        romano = m.group(2).upper()
        anio = m.group(3)
        mes_num = romanos.get(romano)
        if mes_num and mes_num in meses_texto:
            return f"{int(dia)} de {meses_texto[mes_num]} de {anio}"
            
    return fecha_str

def numero_a_letras(numero):
    """Convierte un monto a letras con centavos escritos en texto completo."""
    unidades = ["", "UN", "DOS", "TRES", "CUATRO", "CINCO", "SEIS", "SIETE", "OCHO", "NUEVE"]
    decenas = ["", "DIEZ", "VEINTE", "TREINTA", "CUARENTA", "CINCUENTA", "SESENTA", "SETENTA", "OCHENTA", "NOVENTA"]
    dieces = ["DIEZ", "ONCE", "DOCE", "TRECE", "CATORCE", "QUINCE", "DIECISÉIS", "DIECISIETE", "DIECIOCHO", "DIECINUEVE"]
    centenas = ["", "CIENTO", "DOSCIENTOS", "TRESCIENTOS", "CUATROCIENTOS", "QUINIENTOS", "SEISCIENTOS", "SETECIENTOS", "OCHOCIENTOS", "NOVECIENTOS"]

    def convert_group(n):
        if n == 0: return ""
        if n == 100: return "CIEN"
        c = n // 100
        d = (n % 100) // 10
        u = n % 10
        res = ""
        if c > 0: res += centenas[c] + " "
        if d == 1:
            res += dieces[u] + " "
        else:
            if d > 0:
                res += decenas[d] + (" Y " if u > 0 else " ")
            if u > 0:
                res += unidades[u] + " "
        return res

    try:
        val = float(str(numero).replace('$', '').replace(',', '').strip())
        enteros = int(val)
        centavos = int(round((val - enteros) * 100))
    except Exception:
        return str(numero)

    if enteros == 0:
        texto_enteros = "CERO PESOS"
    else:
        millones = enteros // 1000000
        miles = (enteros % 1000000) // 1000
        cientos = enteros % 1000

        partes = []
        if millones > 0:
            if millones == 1:
                partes.append("UN MILLÓN")
            else:
                partes.append(f"{convert_group(millones).strip()} MILLONES")
        if miles > 0:
            if miles == 1:
                partes.append("MIL")
            else:
                partes.append(f"{convert_group(miles).strip()} MIL")
        if cientos > 0:
            partes.append(convert_group(cientos).strip())

        texto_enteros = " ".join(partes) + " PESOS"

    if centavos > 0:
        texto_centavos = f"CON {convert_group(centavos).strip()} CENTAVOS"
    else:
        texto_centavos = "CON CERO CENTAVOS"

    return f"{texto_enteros} {texto_centavos}, MONEDA NACIONAL"

def extraer_oficina_registral(texto):
    match = re.search(r'OFICINA\s+DE\s+["“\']?([^"”\'\n\r]+?)["”\']?\s+INMUEBLES', texto, re.IGNORECASE)
    if not match:
        match = re.search(r'OFICINA\s+REGISTRAL\s*:\s*([A-ZÁÉÍÓÚÑ\s]+)', texto, re.IGNORECASE)
    if not match:
        match = re.search(r'OFICINA\s+REGISTRAL\s+DE\s+["“\']?([^"”\'\n\r]+?)["”\']?(?=\s+INMUEBLES|\n|,|\.|$)', texto, re.IGNORECASE)
    if match:
        return match.group(1).strip()
    return "NO_ENCONTRADO"

def armar_ubicacion_inmueble(texto_completo):
    match_seccion = re.search(r'DATOS DE IDENTIFICACIÓN[:\s]*(.*?)(?=DATOS DE REGISTRO|INFORMACIÓN COMPLEMENTARIA|VOLANTE|ENTRADA|$)', texto_completo, re.IGNORECASE | re.DOTALL)
    texto_busqueda = match_seccion.group(1) if match_seccion else texto_completo

    patrones_ordenados = [
        ("vivienda", r'(?:Vivienda)[:\s#]*([A-Z0-9\-]+)'),
        ("uso_suelo", r'(?:Uso\s+de?\s+suelo)[:\s#]*([A-ZÁÉÍÓÚ\s]+?)(?=\s*(?:CALLE|LOTE|NO\.|MZ|COL|$))'),
        ("no_interior", r'(?:No\.\s*Interior|Num\.\s*Int\.|Int\.)[:\s#]*([A-Z0-9\-]+)'),
        ("lote", r'(?:Lote)[:\s#]*([A-Z0-9\-]+)'),
        ("manzana", r'(?:Manzana|Mz\.)[:\s#]*([A-Z0-9\-]+)'),
        ("supermanzana", r'(?:Supermanzana|Smz\.)[:\s#]*([A-Z0-9\-]+)'),
        ("etapa", r'(?:Etapa)[:\s#]*([A-Z0-9\-]+)(?!\s*FICIE)'),
        ("condominio", r'(?:Condominio)[:\s#]*([A-Z0-9\-\s]+?)(?=\s*(?:CALLE|NO\.|COL|SECTOR|$))'),
        ("calle", r'(?:Calle|Andador|Avenida|Av\.)[:\s#]*([A-ZÁÉÍÓÚÑ0-9\s]+?)(?=\s*(?:NO\.|NUM\.|LOTE|MZ|COL|C\.P\.|\d|$))'),
        ("no_exterior", r'(?:No\.\s*Exterior|Num\.\s*Ext\.|Ext\.)[:\s#]*([A-Z0-9\-]+)'),
        ("denominacion", r'(?:Denominación\s+del\s+Inmueble|Conjunto|Fraccionamiento)[:\s#]*([A-ZÁÉÍÓÚÑ0-9\s]+?)(?=\s*(?:COLONIA|SECTOR|MUNICIPIO|$))'),
        ("seccion", r'(?:Sección)[:\s#]*([A-Z0-9\-]+)'),
        ("colonia", r'(?:Colonia|Col\.)[:\s#]*([A-ZÁÉÍÓÚÑ0-9\s]+?)(?=\s*(?:SECTOR|MUNICIPIO|C\.P\.|$))'),
        ("sector", r'(?:Sector)[:\s#]*([A-Z0-9\-]+)'),
        ("municipio", r'(?:Municipio|Alcaldía)[:\s#]*([A-ZÁÉÍÓÚÑ\s]+?)(?=\s*(?:DISTRITO|ESTADO|C\.P\.|$))'),
        ("distrito", r'(?:Distrito)[:\s#]*([A-ZÁÉÍÓÚÑ\s]+)'),
        ("estado", r'(?:Estado|Entidad\s+Federativa)[:\s#]*([A-ZÁÉÍÓÚÑ\s]+?)(?=\s*(?:OBSERVACIONES|SUPERFICIE|C\.P\.|$))'),
        ("observaciones", r'(?:Observaciones)[:\s#]*([A-ZÁÉÍÓÚÑ0-9\s]+?)(?=\s*(?:C\.P\.|SUPERFICIE|VOLANTE|$))'),
        ("codigo_postal", r'(?:C\.P\.|Código\s+Postal)[:\s#]*(\d{5})')
    ]

    componentes_encontrados = []
    for clave, patron in patrones_ordenados:
        coincidencia = re.search(patron, texto_busqueda, re.IGNORECASE)
        if coincidencia:
            valor = coincidencia.group(1).strip()
            valor_limpio = re.sub(r'^(CALLE|LOTE|MANZANA|COLONIA|ESTADO)\b', '', valor, flags=re.IGNORECASE).strip()
            if valor_limpio and valor_limpio.lower() not in ['ninguna', 'ninguno', 'no_encontrado', 'no consta', 'null', 'super']:
                etiqueta = clave.replace('_', ' ').title()
                componentes_encontrados.append(f"{etiqueta}: {valor_limpio}")

    if componentes_encontrados:
        return ", ".join(componentes_encontrados)
    return "NO_ENCONTRADO"

def determinar_genero_y_estado_civil(texto_completo, nombre_acreditado=""):
    texto_upper = texto_completo.upper()
    nombre_upper = nombre_acreditado.upper()
    
    if re.search(r'\bLA\s+ACREDITADA\b|\bSEÑORA\b|\bA\s+FAVOR\s+DE\s+LA\b|\bCIUDADANA\b', texto_upper):
        genero = "FEMENINO"
    elif re.search(r'\bEL\s+ACREDITADO\b|\bSEÑOR\b|\bA\s+FAVOR\s+DEL\b|\bCIUDADANO\b', texto_upper):
        genero = "MASCULINO"
    else:
        primer_nombre = nombre_upper.split()[0] if nombre_upper else ""
        if primer_nombre.endswith(('A', 'IA', 'IS')):
            genero = "FEMENINO"
        else:
            genero = "MASCULINO"

    if re.search(r'\bCASAD[AO]\b|\bSOCIEDAD\s+CONYUGAL\b|\bSEPARACION\s+DE\s+BIENES\b|\bESTADO\s+CIVIL\s*:\s*CASAD', texto_upper):
        estado_civil = "CASADO"
    elif re.search(r'\bSOLTER[AO]\b|\bESTADO\s+CIVIL\s*:\s*SOLTER', texto_upper):
        estado_civil = "SOLTERO"
    else:
        estado_civil = "SOLTERO"

    return genero, estado_civil

def extraer_datos_pdf(ruta_pdf):
    datos = {
        "numero_carta": "NO_ENCONTRADO",
        "fecha_expedicion": "NO_ENCONTRADO",
        "numero_credito": "NO_ENCONTRADO",
        "numero_credito_letras": "NO_ENCONTRADO",
        "nombre_acreditado": "NO_ENCONTRADO",
        "monto_credito": "NO_ENCONTRADO",
        "monto_credito_letras": "NO_ENCONTRADO",
        "credito_a_salario": "NO_ENCONTRADO",
        "entidad_financiera": "NO_ENCONTRADO",
        "fecha_liquidacion": "NO_ENCONTRADO",
        "folio_real": "NO_ENCONTRADO",
        "oficina_registral": "NO_ENCONTRADO",
        "datos_inmueble": "NO_ENCONTRADO",
        "genero": "NO_ENCONTRADO",
        "estado_civil": "NO_ENCONTRADO",
        # 🆕 NUEVOS CAMPOS AGREGADOS:
        "numero_escritura": "NO_ENCONTRADO",
        "fecha_escritura": "NO_ENCONTRADO",
        "notario_origen_completo": "NO_ENCONTRADO",
        "tiene_conyuge": "NO_ENCONTRADO",
    }

    try:
        doc = fitz.open(ruta_pdf)
        texto_completo = ""
        for pagina in doc:
            texto_completo += pagina.get_text() + "\n"
        doc.close()
        datos["texto_raw"] = texto_completo
    except Exception as e:
        print(f"Error al abrir PDF {ruta_pdf}: {e}")
        return datos

    nombre_archivo = os.path.basename(ruta_pdf)
    texto_limpio = " ".join(texto_completo.split())

    # 1. NÚMERO DE CRÉDITO
    match_credito = re.search(r'(?:Crédito|Cuenta|Finiquito|Contrato)[:\s#]*(\d{8,12})', texto_limpio, re.IGNORECASE)
    raw_credito = None
    if match_credito:
        raw_credito = match_credito.group(1).strip()
    else:
        match_nombre = re.search(r'(\d{8,12})', nombre_archivo)
        if match_nombre:
            raw_credito = match_nombre.group(1)

    if raw_credito:
        datos["numero_credito"] = raw_credito
        datos["numero_credito_letras"] = credito_a_letras(raw_credito)

    # 2. NÚMERO DE CARTA Y FECHA DE EXPEDICIÓN
    match_carta = re.search(r'Número\s+de\s+carta[:\s#]*([A-Z0-9\-\/]{5,20})', texto_limpio, re.IGNORECASE)
    if not match_carta:
        match_carta = re.search(r'(?:Carta\s*(?:de\s*Instrucción)?|Oficio|Instrucción|Ref)[:\s\.\°\#-]*([A-Z0-9\-\/]{5,20})', texto_limpio, re.IGNORECASE)
    if match_carta:
        val = match_carta.group(1).strip()
        if val.lower() not in ['de', 'del', 'para', 'con', 'que', 'por', 'ext']:
            datos["numero_carta"] = val

    patrones_fecha_exp = [
        r'(?:Fecha\s+de\s+expedici[óo]n|Expedid[oa]\s+el|M[ée]xico,?\s*(?:D\.?F\.?|CDMX)?,?\s*a|Ciudad\s+de\s+M[ée]xico,?\s*a|A)\s*[:\s]*(\d{1,2}\s+de\s+[a-zA-ZÁÉÍÓÚáéíóú]+\s+de\s+\d{4})',
        r'(\d{1,2}\s+de\s+(?:Enero|Febrero|Marzo|Abril|Mayo|Junio|Julio|Agosto|Septiembre|Octubre|Noviembre|Diciembre)\s+de\s+\d{4})'
    ]
    for patron in patrones_fecha_exp:
        match_fecha_exp = re.search(patron, texto_completo, re.IGNORECASE)
        if match_fecha_exp:
            datos["fecha_expedicion"] = match_fecha_exp.group(1).strip()
            break

    # 3. MONTO DEL CRÉDITO Y CÁLCULO DE CRÉDITO A SALARIO (VSM)
    match_vsm = re.search(r'([\d\.]+)\s*VSM', texto_completo, re.IGNORECASE)
    if match_vsm:
        datos["credito_a_salario"] = match_vsm.group(1).strip()

    match_monto = re.search(r'(?:crédito\s+hasta\s+por\s+la\s+cantidad\s+de|monto\s+del?\s+crédito|suerte\s+principal|importe|monto)[:\s]*\$?\s*([\d,]+\.\d{2})', texto_limpio, re.IGNORECASE)
    monto_raw = None
    if match_monto:
        monto_raw = match_monto.group(1).replace(',', '')
    else:
        monto_gen = re.search(r'\$\s*([\d,]+\.\d{2})', texto_limpio)
        if monto_gen:
            monto_raw = monto_gen.group(1).replace(',', '')

    if monto_raw:
        try:
            num = float(str(monto_raw).replace('$', '').replace(',', '').strip())
            datos["monto_credito"] = f"${num:,.2f}"
            datos["monto_credito_letras"] = numero_a_letras(num)
            
            if datos["credito_a_salario"] == "NO_ENCONTRADO":
                veces_salario = num / SALARIO_MINIMO_MENSUAL_DF
                datos["credito_a_salario"] = f"{veces_salario:.2f}"
        except Exception:
            datos["monto_credito"] = monto_raw
            datos["monto_credito_letras"] = monto_raw

    # 4. ENTIDAD FINANCIERA
    if "INFONAVIT" in texto_limpio.upper() or "FONDO NACIONAL DE LA VIVIENDA" in texto_limpio.upper():
        datos["entidad_financiera"] = "INFONAVIT"
    elif "FOVISSSTE" in texto_limpio.upper():
        datos["entidad_financiera"] = "FOVISSSTE"
    elif "BANAMEX" in texto_limpio.upper() or "CITIBANAMEX" in texto_limpio.upper():
        datos["entidad_financiera"] = "BANAMEX"
    elif "BANORTE" in texto_limpio.upper():
        datos["entidad_financiera"] = "BANORTE"

    # 5. FOLIO REAL ELECTRÓNICO
    match_folio = re.search(r'(?:FOLIO\s+REAL\s+ELECTRÓNICO\s+NUMERO|FOLIO\s+REAL\s+ELECTRÓNICO|FOLIO\s+ELECTRÓNICO)[:\s#]*([0-9A-Z\-]{4,15})', texto_limpio, re.IGNORECASE)
    if not match_folio:
        match_folio = re.search(r'(?:Folio\s*Real|Antecedente|F\.R\.|F\.E\.)[:\s#]*([0-9A-Z\-]{4,15})', texto_limpio, re.IGNORECASE)
    if match_folio:
        val = match_folio.group(1).strip()
        if val.lower() not in ['sreales', 'real', 'registral', 'electronico', 'numero']:
            datos["folio_real"] = limpiar_ceros_izquierda(val)

    # 6. NOMBRE DEL ACREDITADO
    patrones_acreditado = [
        r'(?:trabajador|acreditado|deudor)\s+((?:J\.\s*)?[A-ZÁÉÍÓÚÑ\s]{8,50}?)(?=\s+para|\s+gravando|\s+con|\s+cumpli|\s+ha|\.|\,)',
        r'a\s+favor\s+del?\s+((?:J\.\s*)?[A-ZÁÉÍÓÚÑ\s]{8,50}?)(?=\s+para|\s+gravando|\s+con|\.|\,)',
        r'(?:Acreditado\(a\)|Acreditado|Titular|Cliente)[:\s]+([A-ZÁÉÍÓÚÑ\s]{8,50})(?=\s+(?:y/o|con|S\.A\.|RFC|CURP|Crédito|Fecha|\d))',
        r'(?:Acreditado|Titular)[:\s]*\n+([A-ZÁÉÍÓÚÑ\s]{8,50})'
    ]

    palabras_invalidas = ["TRAMITE", "LIBERACION", "CANCELACION", "INMUEBLE", "CREDITO", "INFONAVIT"]

    for patron in patrones_acreditado:
        coincidencia = re.search(patron, texto_completo, re.IGNORECASE)
        if coincidencia:
            nombre = coincidencia.group(1).strip()
            nombre_limpio = " ".join(nombre.split())
            if not any(palabra in nombre_limpio.upper() for palabra in palabras_invalidas) and len(nombre_limpio) > 5:
                datos["nombre_acreditado"] = nombre_limpio
                break

    # 7. FECHA DE LIQUIDACIÓN / PAGO
    match_fecha_pago = re.search(r'(?:saldo\s+deudor.*?:?|a\s+partir\s+de|liquidad[oa]\s+el|pagad[oa]\s+el|fecha\s+de\s+pago)[:\s]*(\d{1,2}\s+de\s+[a-zA-ZÁÉÍÓÚáéíóú]+\s+de\s+\d{4})', texto_completo, re.IGNORECASE)
    if match_fecha_pago:
        datos["fecha_liquidacion"] = match_fecha_pago.group(1).strip()
    else:
        fechas = re.findall(r'(\d{1,2}\s+de\s+[a-zA-ZÁÉÍÓÚáéíóú]+\s+de\s+\d{4})', texto_completo, re.IGNORECASE)
        if len(fechas) > 1:
            datos["fecha_liquidacion"] = fechas[1].strip()
        elif fechas:
            datos["fecha_liquidacion"] = fechas[0].strip()

    # 8. OFICINA REGISTRAL Y UBICACIÓN DEL INMUEBLE
    datos["oficina_registral"] = extraer_oficina_registral(texto_completo)
    datos["datos_inmueble"] = armar_ubicacion_inmueble(texto_completo)

    # 9. DETERMINAR GÉNERO Y ESTADO CIVIL
    genero, estado_civil = determinar_genero_y_estado_civil(texto_completo, datos.get("nombre_acreditado", ""))
    datos["genero"] = genero
    datos["estado_civil"] = estado_civil

   # ════════════════════════════════════════════════════════════════════════
    # 10. EXTRACCIÓN Y LIMPIEZA DE DATOS DE LA CONSTANCIA (IFREM)
    # ════════════════════════════════════════════════════════════════════════

    # A. Número de Escritura
    match_esc = re.search(r'(?:ESCRITURA|INSTRUMENTO)\s+(?:PÚBLICA\s+)?(?:N[Oº°]|NÚM[EÉ]RO|NUMERO)\.?\s*([\d,\.]+)', texto_limpio, re.IGNORECASE)
    if match_esc:
        datos["numero_escritura"] = match_esc.group(1).strip()

    # B. Fecha de Escritura (Convierte romanos a texto formal)
    match_f_esc = re.search(r'DE\s+FECHA\s+([0-9A-Za-z\-\/]+)', texto_limpio, re.IGNORECASE)
    if match_f_esc:
        fecha_bruta = match_f_esc.group(1).strip()
        datos["fecha_escritura"] = limpiar_fecha_escritura(fecha_bruta)

    # C. Notario Origen Completo (Permite leer el punto de "LIC." y extrae hasta el estado)
    match_not = re.search(r'(?:NOTARIO|NOTIARIO)\s+PÚBLICO\s+([A-ZÁÉÍÓÚÑ\.\s0-9]+?(?:ESTADO\s+DE\s+M[EÉ]XICO|M[EÉ]XICO|MEXICO))', texto_limpio, re.IGNORECASE)
    if match_not:
        datos["notario_origen_completo"] = match_not.group(1).strip()

    # D. Si tiene cónyuge (Sí/No)
    tiene_conyuge_match = bool(re.search(r'(C[ÓO]NYUGE|SOCIEDAD\s+CONYUGAL|CASAD[AO]\s+EN\s+SOCIEDAD|EN\s+COPROPIEDAD)', texto_limpio, re.IGNORECASE))
    datos["tiene_conyuge"] = "Sí" if tiene_conyuge_match else "No"

    return datos


def combinar_datos_pareja(datos_lista):
    datos_finales = {
        "numero_carta": "NO_ENCONTRADO",
        "fecha_expedicion": "NO_ENCONTRADO",
        "numero_credito": "NO_ENCONTRADO",
        "numero_credito_letras": "NO_ENCONTRADO",
        "nombre_acreditado": "NO_ENCONTRADO",
        "monto_credito": "NO_ENCONTRADO",
        "monto_credito_letras": "NO_ENCONTRADO",
        "credito_a_salario": "NO_ENCONTRADO",
        "entidad_financiera": "NO_ENCONTRADO",
        "fecha_liquidacion": "NO_ENCONTRADO",
        "folio_real": "NO_ENCONTRADO",
        "oficina_registral": "NO_ENCONTRADO",
        "datos_inmueble": "NO_ENCONTRADO",
        "numero_escritura": "NO_ENCONTRADO",
        "fecha_escritura": "NO_ENCONTRADO",
        "notario_origen_completo": "NO_ENCONTRADO",
        "tiene_conyuge": "NO_ENCONTRADO"
    }
    
    for d in datos_lista:
        for k, v in d.items():
            val = str(v).strip()
            if val != "NO_ENCONTRADO" and val.lower() not in ['de', 'sreales', 'folio', 'carta', 'real']:
                if datos_finales.get(k) == "NO_ENCONTRADO" or len(val) > len(str(datos_finales.get(k, ""))):
                    datos_finales[k] = val
                
    return datos_finales


def reemplazar_texto_en_parrafo(parrafo, mapa_reemplazos):
    texto_parrafo = parrafo.text
    if not any(key in texto_parrafo for key in mapa_reemplazos.keys()):
        return

    for key in mapa_reemplazos.keys():
        if key in parrafo.text and not any(key in r.text for r in parrafo.runs):
            for idx, run in enumerate(parrafo.runs):
                if "{" in run.text:
                    j = idx + 1
                    while j < len(parrafo.runs) and "}" not in parrafo.runs[j-1].text:
                        run.text += parrafo.runs[j].text
                        parrafo.runs[j].text = ""
                        j += 1

    for key, value in mapa_reemplazos.items():
        val_str = str(value)
        for run in parrafo.runs:
            if key in run.text:
                run.text = run.text.replace(key, val_str)
                run.font.underline = False
                run.font.highlight_color = None


def generar_word_cancelacion(ruta_plantilla, datos, ruta_salida):
    if not os.path.exists(ruta_plantilla):
        ruta_plantilla_alt = os.path.join("templates", "plantilla_manera2.docx")
        if os.path.exists(ruta_plantilla_alt):
            ruta_plantilla = ruta_plantilla_alt
        else:
            print(f"ADVERTENCIA: No se encontró la plantilla en {ruta_plantilla}")
            return False

    doc = Document(ruta_plantilla)
    
    folio_raw = str(datos.get("folio_real", "")).strip()
    folio_limpio = limpiar_ceros_izquierda(folio_raw) if folio_raw and folio_raw != "NO_ENCONTRADO" else ""

    credito_raw = str(datos.get("numero_credito", "")).strip()
    if credito_raw and credito_raw != "NO_ENCONTRADO":
        credito_texto = credito_raw if ('"' in credito_raw or '(' in credito_raw) else credito_a_letras(credito_raw)
    else:
        credito_texto = ""

    monto_raw = str(datos.get("monto_credito", "")).strip()
    monto_letras = str(datos.get("monto_credito_letras", "")).strip()
    
    if monto_raw and monto_raw != "NO_ENCONTRADO":
        if not monto_letras or monto_letras == "NO_ENCONTRADO":
            monto_letras = numero_a_letras(monto_raw)
        monto_texto = f"{monto_raw} ({monto_letras})"
    else:
        monto_texto = ""

    def obtener_valor(clave):
        val = str(datos.get(clave, "")).strip()
        return "" if val == "NO_ENCONTRADO" else val

    # Mapeo completo de variables en la plantilla Word (incluyendo los nuevos campos)
    mapa_reemplazos = {
        "{{ numero_carta }}": obtener_valor("numero_carta"),
        "{{numero_carta}}": obtener_valor("numero_carta"),
        "{numero_carta}": obtener_valor("numero_carta"),

        "{{ fecha_expedicion }}": obtener_valor("fecha_expedicion"),
        "{{fecha_expedicion}}": obtener_valor("fecha_expedicion"),
        "{fecha_expedicion}": obtener_valor("fecha_expedicion"),

        "{{ crédito_a_salario }}": obtener_valor("credito_a_salario"),
        "{{credito_a_salario}}": obtener_valor("credito_a_salario"),
        "{credito_a_salario}": obtener_valor("credito_a_salario"),
        
        "{{ numero_credito }}": credito_texto,
        "{{numero_credito}}": credito_texto,
        "{numero_credito}": credito_texto,

        "{{ nombre_acreditado }}": obtener_valor("nombre_acreditado"),
        "{{nombre_acreditado}}": obtener_valor("nombre_acreditado"),
        "{nombre_acreditado}": obtener_valor("nombre_acreditado"),

        "{{ monto_credito }}": monto_texto,
        "{{monto_credito}}": monto_texto,
        "{monto_credito}": monto_texto,

        "{{ monto_letras }}": monto_letras,
        "{{monto_letras}}": monto_letras,
        "{monto_letras}": monto_letras,

        "{{ entidad_financiera }}": obtener_valor("entidad_financiera"),
        "{{entidad_financiera}}": obtener_valor("entidad_financiera"),
        "{entidad_financiera}": obtener_valor("entidad_financiera"),

        "{{ fecha_liquidacion }}": obtener_valor("fecha_liquidacion"),
        "{{fecha_liquidacion}}": obtener_valor("fecha_liquidacion"),
        "{fecha_liquidacion}": obtener_valor("fecha_liquidacion"),

        "{{ folio_real }}": folio_limpio,
        "{{folio_real}}": folio_limpio,
        "{folio_real}": folio_limpio,

        "{{ oficina_registral }}": obtener_valor("oficina_registral"),
        "{{oficina_registral}}": obtener_valor("oficina_registral"),
        "{oficina_registral}": obtener_valor("oficina_registral"),

        "{{ datos_inmueble }}": obtener_valor("datos_inmueble"),
        "{{datos_inmueble}}": obtener_valor("datos_inmueble"),
        "{datos_inmueble}": obtener_valor("datos_inmueble"),

        # 🆕 Nuevas etiquetas para Word:
        "{{ numero_escritura }}": obtener_valor("numero_escritura"),
        "{{numero_escritura}}": obtener_valor("numero_escritura"),

        "{{ fecha_escritura }}": obtener_valor("fecha_escritura"),
        "{{fecha_escritura}}": obtener_valor("fecha_escritura"),

        "{{ notario_origen_completo }}": obtener_valor("notario_origen_completo"),
        "{{notario_origen_completo}}": obtener_valor("notario_origen_completo"),

        "{{ tiene_conyuge }}": obtener_valor("tiene_conyuge"),
        "{{tiene_conyuge}}": obtener_valor("tiene_conyuge"),
    }

    for p in doc.paragraphs:
        reemplazar_texto_en_parrafo(p, mapa_reemplazos)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    reemplazar_texto_en_parrafo(p, mapa_reemplazos)

    directorio_salida = os.path.dirname(ruta_salida)
    if directorio_salida:
        os.makedirs(directorio_salida, exist_ok=True)
        
    doc.save(ruta_salida)
    return True

    